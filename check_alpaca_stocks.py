
import requests
import openpyxl
import csv
from getpass import getpass

EXCEL_FILE = "Final_Algorithmic_Trading_Portfolio_36_Stocks_Final_WDC.xlsx"
BASE_URL = "https://paper-api.alpaca.markets/v2/assets"

api_key = getpass("Enter your PAPER API Key: ")
api_secret = getpass("Enter your PAPER Secret Key: ")

headers = {
    "APCA-API-KEY-ID": api_key,
    "APCA-API-SECRET-KEY": api_secret,
}

wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
ws = wb["Final Portfolio"]

stocks = []
for row in range(5, 41):
    ticker = ws.cell(row, 2).value
    member = ws.cell(row, 4).value
    if ticker:
        stocks.append((str(ticker).strip().upper(), member))

results = []

for ticker, member in stocks:
    try:
        response = requests.get(
            f"{BASE_URL}/{ticker}",
            headers=headers,
            timeout=20
        )

        if response.status_code == 200:
            asset = response.json()
            status = asset.get("status", "unknown")
            tradable = asset.get("tradable", False)

            if status == "active" and tradable:
                result = "AVAILABLE"
            else:
                result = "NOT CURRENTLY TRADABLE"

            company = asset.get("name", "")
            exchange = asset.get("exchange", "")

        else:
            status = f"HTTP {response.status_code}"
            tradable = ""
            company = ""
            exchange = ""
            result = f"CHECK FAILED: {response.text[:150]}"

        print(f"{ticker}: {result}")

        results.append({
            "Member": member,
            "Ticker": ticker,
            "Company": company,
            "Exchange": exchange,
            "Status": status,
            "Tradable": tradable,
            "Result": result,
        })

    except requests.RequestException as error:
        print(f"{ticker}: CONNECTION ERROR: {error}")
        results.append({
            "Member": member,
            "Ticker": ticker,
            "Company": "",
            "Exchange": "",
            "Status": "connection error",
            "Tradable": "",
            "Result": str(error),
        })

with open(
    "alpaca_tradability_check.csv",
    "w",
    newline="",
    encoding="utf-8"
) as file:
    columns = [
        "Member", "Ticker", "Company", "Exchange",
        "Status", "Tradable", "Result"
    ]
    writer = csv.DictWriter(file, fieldnames=columns)
    writer.writeheader()
    writer.writerows(results)

available = sum(r["Result"] == "AVAILABLE" for r in results)

print("\n========== SUMMARY ==========")
print(f"Stocks checked: {len(results)}")
print(f"Available: {available}")
print(f"Other results: {len(results) - available}")
print("Saved report: alpaca_tradability_check.csv")
